#!/usr/bin/env python3
"""Jednorazovy import Robertova rucniho xlsx exportu ceniku z logiman.cz
(administrace Shoptetu) do `logiman_cz_price_export` - bot16, 2026-09-24
(Robert pres bot3, WORKFLOW.md pravidlo 52 "nic se neodklada").

*** DULEZITE ROZLISENI *** (viz i sql/2026-09-24f_logiman_cz_price_export.sql):
tohle NENI crawl - to uz existuje samostatne (scripts/2026-09-24_logiman_
price_reference_crawl.py -> tabulka logiman_cz_price_reference, pouzita v
zalozce "Ceny profilu"). Tenhle skript jen jednou nacte KONKRETNI soubor,
ktery Robert sam exportoval a dal do repa
(backups/2026-09-24_logiman_cenik_export.xlsx), do JINE tabulky pro
zalozku "Dogus: cena vs. vzorec". `source_label`/`imported_at` u kazdeho
radku tenhle puvod zaznamenavaji primo v datech.

Format xlsx (list "Export produktů", 616 datovych radku, header na radku 1):
  sloupec A = code       -> nase shop_products.sku (parovaci klic, Robert potvrdil)
  sloupec C = name       -> nazev na logiman.cz (muze se lisit od naseho nazvu)
  sloupec D = guid       -> Shoptet interni ID, jen pro dohledani
  sloupec E = price      -> cena v Kc BEZ DPH (sloupec includingVat='0' u vsech radku, overeno rucne)
  ostatni sloupce (B, F-AM) jsou pro tenhle import irelevantni (varianty/
  DPH/akce/sklad - nepouzivame).

Pouziti:
    python3 scripts/2026-09-24f_logiman_price_export_import.py            # dry-run, jen vypis
    python3 scripts/2026-09-24f_logiman_price_export_import.py --apply    # skutecny zapis do DB
"""
import argparse
import datetime
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__))))
from _env import load_env  # noqa: E402

import openpyxl  # noqa: E402
import pymysql  # noqa: E402

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_PATH = os.path.join(REPO_ROOT, "backups", "2026-09-24_logiman_cenik_export.xlsx")
SOURCE_LABEL = "robert_xlsx_export_2026-09-24"
SOURCE_FILE_REL = "backups/2026-09-24_logiman_cenik_export.xlsx"


def read_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    assert header[0] == "code" and header[2] == "name" and header[3] == "guid" and header[4] == "price", (
        f"Neocekavana hlavicka xlsx (A/C/D/E melo byt code/name/guid/price): {header[:5]}"
    )
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = row[0]
        if not code:
            continue
        out.append({
            "sku": str(code).strip(),
            "name": row[2],
            "guid": row[3],
            "price_czk": row[4],
        })
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--apply", action="store_true", help="Skutecne zapsat do DB. Bez teto volby jen dry-run.")
    ap.add_argument("--xlsx", default=XLSX_PATH, help="Cesta k xlsx (default: repo backups/ soubor od Roberta).")
    args = ap.parse_args()

    print(f"Cetu {args.xlsx} ...")
    rows = read_rows(args.xlsx)
    print(f"{len(rows)} radku se SKU (nenulovy sloupec A) k importu.")

    dup_skus = [s for s in {r["sku"] for r in rows} if sum(1 for r2 in rows if r2["sku"] == s) > 1]
    if dup_skus:
        print(f"POZOR: {len(dup_skus)} SKU se v souboru opakuje vicekrat: {dup_skus[:10]}", file=sys.stderr)

    imported_at = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None)

    conn = None
    if args.apply:
        env = load_env()
        conn = pymysql.connect(
            host=env["DB_HOST"], port=int(env.get("DB_PORT", 3306)), user=env["DB_USER"],
            password=env["DB_PASSWORD"], database=env["DB_NAME"], charset="utf8mb4",
            cursorclass=pymysql.cursors.DictCursor,
        )

    n = 0
    for r in rows:
        price = float(r["price_czk"]) if r["price_czk"] is not None else None
        print(f"  {r['sku']:<24} {price if price is not None else '(bez ceny)':<10} {r['name']}")
        if args.apply:
            with conn.cursor() as cur:
                cur.execute(
                    "INSERT INTO logiman_cz_price_export "
                    "(sku, product_name, guid, price_czk, source_label, source_file, imported_at) "
                    "VALUES (%s,%s,%s,%s,%s,%s,%s) "
                    "ON DUPLICATE KEY UPDATE product_name=VALUES(product_name), guid=VALUES(guid), "
                    "price_czk=VALUES(price_czk), source_label=VALUES(source_label), "
                    "source_file=VALUES(source_file), imported_at=VALUES(imported_at)",
                    (r["sku"], r["name"], r["guid"], price, SOURCE_LABEL, SOURCE_FILE_REL, imported_at),
                )
            n += 1
    if args.apply:
        conn.commit()
        conn.close()
    print()
    print(f"Hotovo: {n if args.apply else len(rows)} radku {'zapsano' if args.apply else '(dry-run, nic nezapsano)'}.")
    if not args.apply:
        print("(spust s --apply pro skutecny zapis do DB)", file=sys.stderr)


if __name__ == "__main__":
    main()
